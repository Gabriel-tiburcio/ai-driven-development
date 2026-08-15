# -*- coding: utf-8 -*-
"""Generates docs/AllStay_Guia_do_Usuario.docx — a step-by-step user guide for guests and hotel staff."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = RGBColor(0x0F, 0x5F, 0xAE)
BRAND_DARK = RGBColor(0x0A, 0x3D, 0x73)
GREY = RGBColor(0x55, 0x55, 0x55)
LOGO = r"C:\Users\Gabriel\Documents\mba\tcc\logo.png"

doc = Document()

# ---- base style ----
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

for i in range(1, 4):
    h = doc.styles[f"Heading {i}"]
    h.font.color.rgb = BRAND if i == 1 else BRAND_DARK
    h.font.name = "Calibri"

section = doc.sections[0]
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)


def add_page_break():
    doc.add_page_break()


def add_step(number, title, body):
    p = doc.add_paragraph()
    run = p.add_run(f"{number}. {title}")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = BRAND_DARK
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    body_p = doc.add_paragraph(body)
    body_p.paragraph_format.left_indent = Cm(0.6)


def add_tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    run = p.add_run("Dica: ")
    run.bold = True
    run.font.color.rgb = BRAND
    p.add_run(text)


def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6)


# ============================================================
# CAPA
# ============================================================
doc.add_picture(LOGO, width=Inches(1.4))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(18)
run = title.add_run("AllStay")
run.font.size = Pt(40)
run.font.bold = True
run.font.color.rgb = BRAND

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Guia do Usuário — Passo a Passo")
run.font.size = Pt(18)
run.font.color.rgb = GREY

subtitle2 = doc.add_paragraph()
subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle2.add_run("Concierge digital para hóspedes e equipe do hotel")
run.font.size = Pt(12)
run.italic = True
run.font.color.rgb = GREY

add_page_break()

# ============================================================
# SUMÁRIO
# ============================================================
h = doc.add_heading("Sumário", level=1)

toc_items = [
    "1. O que é o AllStay",
    "2. Guia do Hóspede",
    "   2.1 Como acessar o AllStay",
    "   2.2 Explorando o catálogo de atividades",
    "   2.3 Fazendo uma reserva",
    "   2.4 Consultando e cancelando reservas",
    "3. Guia da Equipe do Hotel (Backoffice)",
    "   3.1 Como fazer login",
    "   3.2 Cadastrando uma nova atividade",
    "   3.3 Adicionando horários e vagas",
    "   3.4 Acompanhando as reservas dos hóspedes",
    "4. Perguntas frequentes",
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(2)

add_page_break()

# ============================================================
# 1. O QUE É O ALLSTAY
# ============================================================
doc.add_heading("1. O que é o AllStay", level=1)
doc.add_paragraph(
    "O AllStay é a plataforma de concierge digital do seu hotel. Em vez de ir até a "
    "recepção para saber o que tem disponível, o hóspede acessa o AllStay pelo celular, "
    "vê tudo o que o hotel oferece — spa, passeios, aulas, restaurantes, eventos — e faz "
    "sua reserva em poucos toques, sem baixar nenhum aplicativo."
)
doc.add_paragraph(
    "Para a equipe do hotel, o AllStay oferece um painel (Backoffice) simples para "
    "cadastrar atividades, definir horários e vagas, e acompanhar as reservas em tempo real."
)
doc.add_paragraph("Este guia está dividido em duas partes:")
add_bullets([
    "Guia do Hóspede — para quem vai usar o AllStay durante a estadia no hotel.",
    "Guia da Equipe do Hotel — para quem vai gerenciar atividades e reservas pelo Backoffice.",
])

add_page_break()

# ============================================================
# 2. GUIA DO HÓSPEDE
# ============================================================
doc.add_heading("2. Guia do Hóspede", level=1)

doc.add_heading("2.1 Como acessar o AllStay", level=2)
add_step(1, "Escaneie o QR code",
    "No seu quarto (ou na recepção), você encontrará um QR code do AllStay. Abra a "
    "câmera do seu celular e aponte para o código — não é necessário baixar nenhum app.")
add_step(2, "Toque no link que aparece",
    "A câmera do celular vai reconhecer o QR code e mostrar um link. Toque nele para "
    "abrir o AllStay diretamente no navegador do seu celular.")
add_step(3, "Pronto — você já está conectado ao seu hotel",
    "O AllStay identifica automaticamente o hotel onde você está hospedado e mostra o "
    "catálogo de atividades. Não é preciso criar conta nem fazer login.")
add_tip(
    "Se preferir, você também pode adicionar o AllStay à tela inicial do celular "
    "(\"Adicionar à tela de início\") para acessá-lo como um aplicativo nas próximas vezes."
)

doc.add_heading("2.2 Explorando o catálogo de atividades", level=2)
add_step(1, "Veja tudo o que o hotel oferece",
    "Na tela principal, você verá cards com as atividades disponíveis: nome, categoria, "
    "duração e preço (ou \"Incluso na estadia\", quando for o caso).")
add_step(2, "Filtre por categoria",
    "Use os botões no topo da tela (Spa, Passeios, Gastronomia, Bem-estar...) para ver "
    "apenas as atividades do seu interesse.")
add_step(3, "Toque em uma atividade para ver detalhes",
    "Toque em qualquer card para abrir a descrição completa da atividade e os horários "
    "disponíveis.")

doc.add_heading("2.3 Fazendo uma reserva", level=2)
add_step(1, "Escolha um horário",
    "Na página da atividade, toque em um dos horários disponíveis. Horários esgotados "
    "aparecem marcados como \"Esgotado\" e não podem ser selecionados.")
add_step(2, "Preencha seus dados",
    "Informe seu nome e o número do seu quarto. Esses dados são usados para identificar "
    "sua reserva e cobrar na conta do quarto, se aplicável.")
add_step(3, "Confirme a reserva",
    "Toque em \"Confirmar reserva\". Você verá uma mensagem de confirmação na tela — "
    "não é necessário pagamento neste momento; o valor (quando houver) é lançado na "
    "conta do quarto e cobrado no checkout.")
add_tip("Chegue com alguns minutos de antecedência ao local da atividade reservada.")

doc.add_heading("2.4 Consultando e cancelando reservas", level=2)
add_step(1, "Acesse \"Minhas reservas\"",
    "No menu inferior do AllStay, toque em \"Minhas reservas\".")
add_step(2, "Informe o número do quarto",
    "Digite o número do seu quarto e toque em \"Buscar\" para ver todas as suas reservas ativas.")
add_step(3, "Cancele se necessário",
    "Caso não possa comparecer, toque em \"Cancelar\" na reserva desejada. A vaga é "
    "liberada imediatamente para outros hóspedes.")

add_page_break()

# ============================================================
# 3. GUIA DA EQUIPE DO HOTEL
# ============================================================
doc.add_heading("3. Guia da Equipe do Hotel (Backoffice)", level=1)
doc.add_paragraph(
    "O Backoffice é o painel administrativo usado pela equipe do hotel para gerenciar o "
    "catálogo de atividades, os horários disponíveis e acompanhar as reservas feitas "
    "pelos hóspedes."
)

doc.add_heading("3.1 Como fazer login", level=2)
add_step(1, "Acesse o endereço do Backoffice",
    "Abra o navegador e acesse o endereço do Backoffice fornecido pela equipe AllStay "
    "(ex.: hotel.allstay.eupanda.com.br).")
add_step(2, "Informe e-mail e senha",
    "Digite o e-mail e a senha cadastrados para sua conta de acesso e toque em \"Entrar\".")
add_step(3, "Você será direcionado à tela de Atividades",
    "Após o login, o sistema mostra automaticamente o catálogo de atividades do seu hotel.")
add_tip(
    "Esqueceu sua senha ou não tem uma conta ainda? Entre em contato com o "
    "administrador do sistema para criar ou redefinir seu acesso."
)

doc.add_heading("3.2 Cadastrando uma nova atividade", level=2)
add_step(1, "Toque em \"+ Nova atividade\"",
    "Na tela de Atividades, toque no botão \"+ Nova atividade\" para abrir o formulário "
    "de cadastro.")
add_step(2, "Preencha as informações",
    "Informe o nome da atividade, a categoria (Spa, Passeios, Gastronomia, etc.), uma "
    "descrição, o preço (ou zero, se for incluso na estadia) e a duração em minutos.")
add_step(3, "Toque em \"Salvar\"",
    "A atividade é criada e passa a aparecer na lista — porém ainda sem horários "
    "disponíveis para reserva (veja o próximo passo).")

doc.add_heading("3.3 Adicionando horários e vagas", level=2)
add_step(1, "Abra \"Gerenciar horários\"",
    "Na lista de atividades, toque em \"Gerenciar horários\" ao lado da atividade desejada.")
add_step(2, "Cadastre um novo horário",
    "Escolha a data e o horário de início, defina a capacidade (número de vagas) e "
    "toque em \"Adicionar horário\".")
add_step(3, "Repita para outros dias/horários",
    "Você pode cadastrar quantos horários forem necessários para a mesma atividade — "
    "por exemplo, um horário por dia durante a semana.")
add_tip(
    "Para remover um horário (por exemplo, em caso de manutenção ou feriado), use o "
    "botão \"Remover\" na lista de horários da atividade."
)

doc.add_heading("3.4 Acompanhando as reservas dos hóspedes", level=2)
add_step(1, "Acesse o menu \"Reservas\"",
    "No topo do Backoffice, toque em \"Reservas\" para ver todas as reservas feitas "
    "pelos hóspedes do hotel.")
add_step(2, "Consulte os detalhes",
    "A lista mostra a atividade reservada, o horário, o nome do hóspede, o número do "
    "quarto e o status da reserva (Confirmada ou Cancelada).")
add_step(3, "Use as informações para se preparar",
    "Utilize essa lista para organizar a equipe de cada atividade com antecedência — "
    "quantas pessoas confirmadas, em qual horário.")

add_page_break()

# ============================================================
# 4. FAQ
# ============================================================
doc.add_heading("4. Perguntas frequentes", level=1)

faqs = [
    ("Preciso instalar algum aplicativo?",
     "Não. O AllStay funciona diretamente no navegador do celular, tanto para hóspedes "
     "quanto para a equipe do hotel."),
    ("O hóspede precisa criar uma conta?",
     "Não. O acesso do hóspede é feito apenas com o número do quarto — sem senha, sem "
     "cadastro."),
    ("Como o pagamento é feito?",
     "Nesta versão, as reservas são lançadas na conta do quarto e cobradas no checkout, "
     "junto com as demais despesas da estadia."),
    ("Posso reservar mais de uma atividade?",
     "Sim. Não há limite de reservas — o hóspede pode reservar quantas atividades "
     "quiser, respeitando a disponibilidade de vagas."),
    ("O que acontece se eu cancelar uma reserva?",
     "A vaga é liberada imediatamente e volta a ficar disponível para outros hóspedes."),
]

for question, answer in faqs:
    p = doc.add_paragraph()
    run = p.add_run(question)
    run.bold = True
    run.font.color.rgb = BRAND_DARK
    doc.add_paragraph(answer)

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("AllStay — Concierge digital para hotéis 4 e 5 estrelas")
run.italic = True
run.font.color.rgb = GREY
run.font.size = Pt(9)

out_path = r"C:\Users\Gabriel\Documents\mba\tcc\docs\AllStay_Guia_do_Usuario.docx"
doc.save(out_path)
print("saved:", out_path)
